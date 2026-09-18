import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def create_document():
    doc = Document()

    # Set page margins (A4)
    sections = doc.sections
    for section in sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.65)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

    # Styles
    # Primary colors (Blue theme)
    PRIMARY_COLOR = RGBColor(37, 99, 235)       # #2563eb Blue
    PRIMARY_DARK = RGBColor(30, 64, 175)       # #1e40af Dark Blue
    TEXT_MAIN = RGBColor(15, 23, 42)           # #0f172a Slate 900
    TEXT_MUTED = RGBColor(71, 85, 105)         # #475569 Slate 600
    BG_LIGHT = "F0F7FF"                        # Light blue fill
    BORDER_COLOR = "CBD5E1"                    # Light gray border

    # Header / Title
    p_badge = doc.add_paragraph()
    r_badge = p_badge.add_run("DARUKAA.EARTH AI BIODIVERSITY INTELLIGENCE CHALLENGE")
    r_badge.font.name = "Arial"
    r_badge.font.size = Pt(8.5)
    r_badge.font.bold = True
    r_badge.font.color.rgb = PRIMARY_DARK

    p_title = doc.add_paragraph()
    r_title = p_title.add_run("EcoSynapse AI: Evidence-Grounded Ecological Intelligence Platform")
    r_title.font.name = "Arial"
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = TEXT_MAIN
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Official Engineering Submission & Technical Specification Document")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(10)
    r_sub.font.bold = True
    r_sub.font.color.rgb = PRIMARY_COLOR
    p_sub.paragraph_format.space_after = Pt(12)

    # Section 1: Project Submission Links
    h1 = doc.add_paragraph()
    r1 = h1.add_run("1. Project Submission Links")
    r1.font.name = "Arial"
    r1.font.size = Pt(12)
    r1.font.bold = True
    r1.font.color.rgb = PRIMARY_DARK
    h1.paragraph_format.space_before = Pt(8)
    h1.paragraph_format.space_after = Pt(4)

    links = [
        ("GitHub Repository:", " https://github.com/daanialmirza5/EcoSynapse-AI"),
        ("Live Cloud Demo (Frontend):", " https://ecosynapse-ai-1.onrender.com")
    ]
    for label, val in links:
        p = doc.add_paragraph(style='List Bullet')
        r_lbl = p.add_run(label)
        r_lbl.bold = True
        r_lbl.font.size = Pt(9.5)
        r_val = p.add_run(val)
        r_val.font.size = Pt(9.5)
        r_val.font.color.rgb = PRIMARY_COLOR
        p.paragraph_format.space_after = Pt(2)

    # Callout box for repo notes
    callout_tbl = doc.add_table(rows=1, cols=1)
    callout_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    callout_cell = callout_tbl.cell(0, 0)
    callout_cell.width = Inches(6.9)
    set_cell_background(callout_cell, "EFF6FF")
    set_cell_margins(callout_cell, top=120, bottom=120, left=180, right=180)
    p_c = callout_cell.paragraphs[0]
    rc_title = p_c.add_run("Repository Access Notes:\n")
    rc_title.bold = True
    rc_title.font.size = Pt(9)
    rc_title.font.color.rgb = PRIMARY_DARK
    rc_body = p_c.add_run("The GitHub repository is hosted at daanialmirza5/EcoSynapse-AI (branch: main). Evaluators can review the complete Git commit history, automated CI logs, verified test suites (79 passing tests), and reproduction scripts.")
    rc_body.font.size = Pt(8.5)
    rc_body.font.color.rgb = TEXT_MUTED

    # Section 2: System Architecture
    h2 = doc.add_paragraph()
    r2 = h2.add_run("2. System Architecture & Information Pipeline")
    r2.font.name = "Arial"
    r2.font.size = Pt(12)
    r2.font.bold = True
    r2.font.color.rgb = PRIMARY_DARK
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(4)

    p_arch = doc.add_paragraph()
    r_arch = p_arch.add_run(
        "EcoSynapse AI operates on an uncompromising design principle: an LLM is NOT the reasoning engine. "
        "Traditional generative AI chatbots suffer from severe hallucinations when handling delicate ecological metrics, "
        "prescribing unconstrained interventions with fabricated citations. EcoSynapse AI replaces ungrounded generative speculation "
        "with a 100% deterministic Python reasoning pipeline, grounded against an auditable scientific corpus and a typed biological knowledge graph."
    )
    r_arch.font.size = Pt(9.5)
    p_arch.paragraph_format.space_after = Pt(6)

    # Insert Architecture Diagram Image
    img_path = os.path.join(os.path.dirname(__file__), "architecture.jpg")
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(4)
        p_img.paragraph_format.space_after = Pt(6)
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=Inches(6.8))

    p_pipe = doc.add_paragraph()
    r_pipe = p_pipe.add_run("The information pipeline processes ecological observations through five connected layers:")
    r_pipe.font.size = Pt(9.5)
    p_pipe.paragraph_format.space_after = Pt(3)

    steps = [
        ("1. User Input & Extraction:", " Natural language text or structured JSON is parsed by a deterministic regex extractor (app/conversations/extraction.py). Gaps trigger targeted clarification questions (app/conversations/clarify.py)."),
        ("2. Hybrid Scientific Retrieval & Knowledge Graph:", " Cosine similarity over embeddings is combined with BM25 lexical keyword search and 1-hop neighborhood expansion across a 23-node, 22-edge typed NetworkX MultiDiGraph."),
        ("3. 10-Step Deterministic Reasoning & Constraint Engine:", " Identifies active concerns, maps causal pathways, and applies strict physical feasibility rules (e.g., rainfall < 500mm limits woody vegetation)."),
        ("4. Scientific Claim Verification:", " Validates every quantitative assertion against raw source excerpts, checking exact numeric match before rating evidence strength."),
        ("5. Structured Output & Monitoring:", " Synthesizes actionable recommendations tagged with time horizons, confidence scores, and metric-specific monitoring cadences (no fabricated baselines).")
    ]
    for step_num, step_desc in steps:
        p = doc.add_paragraph()
        r_s = p.add_run(step_num)
        r_s.bold = True
        r_s.font.size = Pt(9)
        r_s.font.color.rgb = PRIMARY_DARK
        r_d = p.add_run(step_desc)
        r_d.font.size = Pt(9)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.space_after = Pt(2)

    doc.add_page_break()

    # Section 3: Core Ecological Domains
    h3 = doc.add_paragraph()
    r3 = h3.add_run("3. Core Ecological Intelligence Domains")
    r3.font.name = "Arial"
    r3.font.size = Pt(12)
    r3.font.bold = True
    r3.font.color.rgb = PRIMARY_DARK
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)

    p_dom = doc.add_paragraph()
    r_dom = p_dom.add_run("EcoSynapse AI models interconnected biological relationships across all five mandatory challenge domains without relying on ungrounded generative speculation:")
    r_dom.font.size = Pt(9.5)
    p_dom.paragraph_format.space_after = Pt(6)

    # Domain Table
    dom_data = [
        ("Domain", "Modeled Variables & Parameters", "Ecological Significance & Mechanisms"),
        ("1. Soil Health", "pH, Organic Matter (SOM %), Bulk Density (g/cm³), N-P-K, Microbial Activity", "Regulates nutrient availability, moisture retention, and structural stability against erosion."),
        ("2. Water Quality", "Dissolved Oxygen (DO mg/L), Turbidity (NTU), Nitrate (mg/L), pH, Flow Velocity", "Drives aquatic biodiversity, bioindicator health (ephemeroptera), and eutrophication risk."),
        ("3. Forest & Vegetation", "Canopy Cover (%), Native vs. Invasive Ratio, Basal Area (m²/ha), DBH Distribution", "Establishes multi-tier vertical strata, carbon sequestration, and seed bank dynamics."),
        ("4. Wildlife & Fauna", "Shannon Diversity Index (H'), Indicator Species Richness, Functional Guilds", "Signals trophic web integrity, apex predator presence, and pollinator mutualisms."),
        ("5. Climate & Atmosphere", "Mean Annual Precipitation (MAP mm), Temp Extremes (°C), VPD (kPa), Albedo", "Imposes biophysical boundary limits on plant physiology, drought resilience, and fire regimes.")
    ]

    dom_tbl = doc.add_table(rows=len(dom_data), cols=3)
    dom_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(dom_tbl.rows):
        is_header = (r_idx == 0)
        for c_idx, cell in enumerate(row.cells):
            cell.paragraphs[0].text = dom_data[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.runs[0]
            r.font.name = "Arial"
            r.font.size = Pt(8.5)
            if is_header:
                r.bold = True
                r.font.color.rgb = PRIMARY_DARK
                set_cell_background(cell, BG_LIGHT)
            else:
                if c_idx == 0:
                    r.bold = True
                r.font.color.rgb = TEXT_MAIN
                if r_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

    # Section 4: Key Technical Differentiators
    h4 = doc.add_paragraph()
    r4 = h4.add_run("4. Key Technical Differentiators & Competitive Benchmarks")
    r4.font.name = "Arial"
    r4.font.size = Pt(12)
    r4.font.bold = True
    r4.font.color.rgb = PRIMARY_DARK
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(4)

    comp_data = [
        ("Evaluation Dimension", "Standard LLM / Chatbots", "EcoSynapse AI Platform"),
        ("Reasoning Engine", "Probabilistic next-token prediction (high hallucination rate on numbers)", "100% Deterministic Python Pipeline (zero metric hallucination)"),
        ("Knowledge Graph", "None or unstructured text embeddings only", "23-node, 22-edge typed NetworkX MultiDiGraph for causal tracing"),
        ("Source Verification", "Unchecked or simulated URL references", "Exact numeric string matching against raw scientific corpus"),
        ("Constraint Checking", "Passive prompting (frequently ignored under complex inputs)", "Hard threshold validators (rainfall, soil pH, moisture limits)"),
        ("Uncertainty Handling", "Masks ignorance with confident assertions", "Explicitly scores confidence (0.0-1.0) and asks targeted clarifications"),
        ("Safety Guarantees", "May recommend invasive species or harmful water alterations", "Deterministic rules block unfeasible or ecologically dangerous interventions"),
        ("Test Suite", "Ad-hoc manual prompting", "79 Automated Tests (61 Backend Pytest + 18 Frontend Vitest)")
    ]

    comp_tbl = doc.add_table(rows=len(comp_data), cols=3)
    comp_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(comp_tbl.rows):
        is_header = (r_idx == 0)
        for c_idx, cell in enumerate(row.cells):
            cell.paragraphs[0].text = comp_data[r_idx][c_idx]
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.runs[0]
            r.font.name = "Arial"
            r.font.size = Pt(8.2)
            if is_header:
                r.bold = True
                r.font.color.rgb = PRIMARY_DARK
                set_cell_background(cell, BG_LIGHT)
            else:
                if c_idx == 0:
                    r.bold = True
                elif c_idx == 2:
                    r.bold = True
                    r.font.color.rgb = PRIMARY_DARK
                else:
                    r.font.color.rgb = TEXT_MUTED
                if r_idx % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)

    # Section 5: Reproduction & Verification Guide
    h5 = doc.add_paragraph()
    r5 = h5.add_run("5. Reproducibility & Local Verification Guide")
    r5.font.name = "Arial"
    r5.font.size = Pt(12)
    r5.font.bold = True
    r5.font.color.rgb = PRIMARY_DARK
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(4)

    p_rep = doc.add_paragraph()
    r_rep = p_rep.add_run("Evaluators can reproduce all results and run the comprehensive test suites locally with standard open-source tools:")
    r_rep.font.size = Pt(9.5)
    p_rep.paragraph_format.space_after = Pt(4)

    commands = (
        "# 1. Clone the repository\n"
        "git clone https://github.com/daanialmirza5/EcoSynapse-AI.git\n"
        "cd EcoSynapse-AI\n\n"
        "# 2. Run Backend Test Suite (61 tests)\n"
        "cd apps/api\n"
        "pytest tests/ -v\n\n"
        "# 3. Run Frontend Test Suite (18 tests)\n"
        "cd ../web\n"
        "npm test -- --run\n\n"
        "# 4. Validate Knowledge Graph & Scientific Corpus\n"
        "python scripts/validate_knowledge.py"
    )
    p_cmd = doc.add_paragraph()
    r_cmd = p_cmd.add_run(commands)
    r_cmd.font.name = "Courier New"
    r_cmd.font.size = Pt(8.5)
    r_cmd.font.color.rgb = RGBColor(226, 232, 240)
    
    # We can wrap command in a table with dark background
    cmd_tbl = doc.add_table(rows=1, cols=1)
    cmd_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    cmd_cell = cmd_tbl.cell(0, 0)
    cmd_cell.width = Inches(6.9)
    set_cell_background(cmd_cell, "0F172A")
    set_cell_margins(cmd_cell, top=100, bottom=100, left=150, right=150)
    p_code = cmd_cell.paragraphs[0]
    r_c = p_code.add_run(commands)
    r_c.font.name = "Courier New"
    r_c.font.size = Pt(8)
    r_c.font.color.rgb = RGBColor(226, 232, 240)

    # Save document
    out_path = os.path.join(os.path.dirname(__file__), "EcoSynapse_AI_Darukaa_Submission.docx")
    doc.save(out_path)
    print("DOCX successfully generated at:", out_path)

if __name__ == "__main__":
    create_document()
